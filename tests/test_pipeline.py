"""Unit tests for extraction + chunking (no vector store needed)."""

import csv

import pytest

from src.ingestion.pipeline import (
    CHUNK_SIZE,
    UnsupportedDocumentError,
    ingest_file,
    read_text_file,
    split_text,
)


def test_split_text_respects_chunk_size_and_keeps_all_words():
    text = " ".join(f"word{i}" for i in range(2000))
    chunks = split_text(text)
    assert len(chunks) > 1
    assert all(len(c) <= CHUNK_SIZE + 150 for c in chunks)
    joined = " ".join(chunks)
    assert "word0" in joined and "word1999" in joined


def test_txt_decodes_non_utf8_instead_of_crashing(tmp_path):
    path = tmp_path / "latin.txt"
    path.write_bytes("Café plant – pH 7.2 at Müller works".encode("cp1252"))
    assert "Café" in read_text_file(path)
    chunks = ingest_file(path)
    assert chunks[0].metadata["file_type"] == "txt"


def test_long_txt_is_split_into_multiple_chunks(tmp_path):
    path = tmp_path / "long.txt"
    path.write_text("\n\n".join(f"Paragraph {i}. " + "filler text " * 40 for i in range(30)))
    chunks = ingest_file(path)
    assert len(chunks) > 5
    assert {c.metadata["content_type"] for c in chunks} == {"text"}
    assert [c.metadata["chunk_index"] for c in chunks] == list(range(len(chunks)))


def test_csv_becomes_labelled_rows_not_one_blob(tmp_path):
    path = tmp_path / "plants.csv"
    with path.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Plant", "Capacity MLD", "Technology"])
        for i in range(50):
            w.writerow([f"Plant{i}", i + 1, "SBR"])
    chunks = ingest_file(path)
    kinds = [c.metadata["content_type"] for c in chunks]
    assert kinds.count("table_record") == 50
    assert kinds.count("table_summary") == 1
    # Never the legacy STP type: that would pollute the STP aggregations.
    assert "table_row" not in kinds
    row = next(c for c in chunks if "Plant7" in c.page_content and c.metadata["content_type"] == "table_record")
    assert "Capacity MLD: 8" in row.page_content


def test_generic_xlsx_with_arbitrary_layout(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Pumps"
    ws.append(["Quarterly pump register"])        # title row above the header
    ws.append([])
    ws.append(["Pump tag", "Duty", "Installed"])
    ws.append(["P-901", "Sludge transfer", "2029-04-01"])
    ws.append(["P-902", "Raw sewage", "2030-01-15"])
    path = tmp_path / "pumps.xlsx"
    wb.save(path)

    chunks = ingest_file(path)
    text = "\n".join(c.page_content for c in chunks)
    assert "Pump tag: P-902" in text and "Duty: Raw sewage" in text
    assert {c.metadata["file_type"] for c in chunks} == {"spreadsheet"}


def test_docx_paragraphs_and_tables(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_heading("Inspection note", level=1)
    doc.add_paragraph("The clarifier weir at Quillon works was re-levelled in June.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Asset"
    table.rows[0].cells[1].text = "Condition"
    table.rows[1].cells[0].text = "Blower B-7"
    table.rows[1].cells[1].text = "Needs bearing replacement"
    path = tmp_path / "note.docx"
    doc.save(path)

    chunks = ingest_file(path)
    kinds = {c.metadata["content_type"] for c in chunks}
    assert {"text", "table_record"} <= kinds
    assert any("re-levelled in June" in c.page_content for c in chunks)
    assert any("Condition: Needs bearing replacement" in c.page_content for c in chunks)


@pytest.mark.parametrize("name,content", [("empty.txt", ""), ("blank.csv", "\n\n")])
def test_empty_files_raise_instead_of_reporting_success(tmp_path, name, content):
    path = tmp_path / name
    path.write_text(content)
    with pytest.raises(UnsupportedDocumentError):
        ingest_file(path)


def test_unsupported_extension_raises(tmp_path):
    path = tmp_path / "legacy.xls"
    path.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(UnsupportedDocumentError):
        ingest_file(path)


def test_text_pdf_narrative_is_chunked(tmp_path):
    reportlab = pytest.importorskip("reportlab")
    from reportlab.pdfgen import canvas

    path = tmp_path / "narrative.pdf"
    c = canvas.Canvas(str(path))
    c.drawString(72, 750, "The Brindavar facility recycles 63 percent of its treated effluent.")
    c.save()

    chunks = ingest_file(path)
    assert any("Brindavar" in ch.page_content for ch in chunks)
    assert chunks[0].metadata["page"] == 1


def test_scanned_pdf_is_ocrd_or_rejected_clearly(tmp_path):
    PIL = pytest.importorskip("PIL")
    pytest.importorskip("pytesseract")
    import shutil

    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1400, 300), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 56)
    except OSError:
        font = ImageFont.load_default()
    draw.text((40, 100), "Kestrel plant discharge permit 4471", fill="black", font=font)
    path = tmp_path / "scan.pdf"
    image.save(path, "PDF")   # image-only PDF: no text layer

    if shutil.which("tesseract"):
        chunks = ingest_file(path)
        assert any("4471" in c.page_content for c in chunks)
        assert chunks[0].metadata.get("ocr") is True
    else:
        with pytest.raises(UnsupportedDocumentError, match="OCR"):
            ingest_file(path)

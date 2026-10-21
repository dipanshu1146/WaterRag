"""End-to-end acceptance tests through the real HTTP routes.

Each document type is uploaded through /api/documents/upload, must report a
positive record count, and must then be retrievable by a natural question that
mentions neither the filename nor the file type.
"""

import io

import pytest


def upload(client, name, data, content_type="application/octet-stream"):
    return client.post("/api/documents/upload", files={"file": (name, io.BytesIO(data), content_type)})


def ask(client, question):
    response = client.post("/api/ask", json={"question": question})
    assert response.status_code == 200, response.text
    return response.json()


def evidence(body):
    return body["answer"] + "\n" + "\n".join(s["content"] for s in body["sources"])


def assert_indexed(response):
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "indexed"
    assert body["indexed_records"] > 0
    return body


# --------------------------------------------------------------------------
# Brand-new documents, one unique fact each
# --------------------------------------------------------------------------

def test_new_txt_with_free_form_paragraphs(client):
    text = (
        "Site memo.\n\nThe Zorblax pumping station was commissioned in March 2031 "
        "and is fitted with vortex aerators supplied by Halvorsen Marine.\n\n"
        "Operators rotate every two weeks."
    )
    assert_indexed(upload(client, "memo.txt", text.encode()))
    body = ask(client, "Who supplied the aerators at Zorblax?")
    assert "Halvorsen Marine" in evidence(body)
    assert body["sources"][0]["document"] == "memo.txt"


def test_new_csv(client):
    csv_bytes = (
        "Plant,District,Effluent reuse,Technology\n"
        "Quillon works,Nilgiri,Irrigation of 14 hectares,Extended aeration\n"
        "Tarn Road,Kolar,Cooling tower makeup,MBBR\n"
    ).encode()
    assert_indexed(upload(client, "sites.csv", csv_bytes))
    body = ask(client, "Where does Tarn Road send its effluent?")
    assert "Cooling tower makeup" in evidence(body)


def test_new_docx_paragraph_and_table(client):
    from docx import Document

    doc = Document()
    doc.add_paragraph("The Ostrava clarifier weir was re-levelled during the June shutdown.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Asset", "Condition"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "Blower B-7", "Needs bearing replacement"
    buffer = io.BytesIO()
    doc.save(buffer)

    assert_indexed(upload(client, "inspection.docx", buffer.getvalue()))
    assert "June shutdown" in evidence(ask(client, "When was the Ostrava weir re-levelled?"))
    assert "bearing replacement" in evidence(ask(client, "What is wrong with blower B-7?"))


def test_new_xlsx_with_arbitrary_layout(client):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Pumps"
    ws.append(["Pump tag", "Duty", "Installed"])
    ws.append(["P-901", "Sludge transfer", "2029-04-01"])
    ws.append(["P-902", "Raw sewage lifting", "2030-01-15"])
    buffer = io.BytesIO()
    wb.save(buffer)

    assert_indexed(upload(client, "register.xlsx", buffer.getvalue()))
    assert "Raw sewage lifting" in evidence(ask(client, "What is pump P-902 used for?"))


def test_new_pdf_narrative(client):
    pytest.importorskip("reportlab")
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer)
    c.drawString(72, 750, "The Brindavar facility recycles 63 percent of its treated effluent.")
    c.save()

    assert_indexed(upload(client, "annual.pdf", buffer.getvalue()))
    assert "63 percent" in evidence(ask(client, "How much effluent does Brindavar recycle?"))


# --------------------------------------------------------------------------
# Upload lifecycle
# --------------------------------------------------------------------------

def test_unrelated_question_gets_no_sources(client):
    assert_indexed(upload(client, "memo.txt", b"The Zorblax pumping station has vortex aerators."))
    body = ask(client, "Who painted the Mona Lisa?")
    assert body["sources"] == []


def test_empty_upload_is_rejected_not_reported_as_indexed(client):
    response = upload(client, "empty.txt", b"   \n")
    assert response.status_code == 422
    assert client.get("/api/documents").json()["total"] == 0


def test_xls_is_rejected_with_clear_message(client):
    response = upload(client, "legacy.xls", b"\xd0\xcf\x11\xe0")
    assert response.status_code == 400


def test_duplicate_upload_is_rejected(client):
    assert_indexed(upload(client, "a.txt", b"Unique duplicate-check content about Zorblax."))
    assert upload(client, "b.txt", b"Unique duplicate-check content about Zorblax.").status_code == 409


def test_failed_indexing_can_be_retried(client, monkeypatch):
    import src.api.routes.upload as upload_module

    real = upload_module.index_uploaded_file

    def boom(path):
        raise RuntimeError("chroma exploded")

    monkeypatch.setattr(upload_module, "index_uploaded_file", boom)
    response = upload(client, "memo.txt", b"The Zorblax pumping station has vortex aerators.")
    assert response.status_code == 500

    listing = client.get("/api/documents").json()["documents"]
    assert listing[0]["status"] == "saved, not indexed"
    stored = listing[0]["stored_filename"]

    monkeypatch.setattr(upload_module, "index_uploaded_file", real)
    retry = client.post(f"/api/documents/{stored}/reindex")
    assert retry.status_code == 200 and retry.json()["indexed_records"] > 0
    assert client.get("/api/documents").json()["documents"][0]["status"] == "indexed"


def test_reindex_does_not_duplicate_records(client, store):
    body = assert_indexed(upload(client, "memo.txt", b"The Zorblax pumping station has vortex aerators."))
    before = len(store.rows)
    assert client.post(f"/api/documents/{body['stored_filename']}/reindex").status_code == 200
    assert len(store.rows) == before


def test_delete_via_api_removes_the_files_records(client, store):
    body = assert_indexed(upload(client, "memo.txt", b"The Zorblax pumping station has vortex aerators."))
    assert len(store.rows) > 0

    response = client.delete(f"/api/documents/{body['stored_filename']}")

    assert response.status_code == 200
    assert response.json()["deleted_indexed_records"] > 0
    assert store.rows == {}
    assert client.get("/api/documents").json()["total"] == 0


def test_delete_matches_relative_and_absolute_source_paths(store):
    """Bulk ingestion stored relative paths; uploads store absolute ones."""
    from src.vectorstore.search_index import PROJECT_ROOT, delete_source_records, source_variants

    target = PROJECT_ROOT / "data" / "uploads" / "report.pdf"
    spellings = [str(target), "data/uploads/report.pdf", "./data/uploads/report.pdf"]
    for number, source in enumerate(spellings):
        store.rows[f"r{number}"] = {"content": "x", "metadata": {"source": source}}
    store.rows["other"] = {"content": "y", "metadata": {"source": "data/uploads/other.pdf"}}

    assert set(spellings) <= set(source_variants(target))
    assert delete_source_records(target) == 3
    assert list(store.rows) == ["other"]


# --------------------------------------------------------------------------
# Regression: the malformed STP-ID regexes used to return HTTP 500
# --------------------------------------------------------------------------

def _seed_stp_row(store):
    store.rows["stp-1"] = {
        "content": (
            "STP ID: YMN-STP-001\n"
            "STP name: STP's Name: Rohtak 18.5 MLD (SBR)\n"
            "STP's Name: Rohtak 18.5 MLD (SBR)\n"
            "Sample type: outlet\n"
        ),
        "metadata": {
            "source": "data/uploads/stp.xlsx", "file_type": "xlsx",
            "content_type": "table_row", "row_number": 7,
            "stp_id": "YMN-STP-001", "sample_type": "outlet",
        },
    }


@pytest.mark.parametrize("question,expected", [
    ("What is the treatment capacity of YMN-STP-001?", "18.5"),
    ("What treatment technology does YMN-STP-001 use?", "SBR"),
])
def test_stp_id_route_no_longer_returns_500(client, store, question, expected):
    _seed_stp_row(store)
    response = client.post("/api/ask", json={"question": question})
    assert response.status_code == 200, response.text
    body = response.json()
    assert expected in body["answer"]
    assert body["records"][0]["capacity_mld"] == 18.5
    assert body["records"][0]["treatment_technology"] == "SBR"

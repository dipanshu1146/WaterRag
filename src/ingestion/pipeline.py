"""Single, authoritative ingestion pipeline for uploaded files.

    extract  ->  chunk (format-aware)  ->  normalise metadata  ->  (index)

This replaces the three competing implementations that previously existed
(src/vectorstore/build_index.py::split_documents,
src/ingestion/chunk_documents.py and src/chunking/chunk_documents.py) for the
live upload path. Those modules are left in place but are no longer wired into
uploads.

Design rules
------------
* Light dependencies only (stdlib, openpyxl, python-docx, pdfplumber). The
  conversion to LangChain ``Document`` objects happens in the indexer, so this
  module can be unit-tested without Chroma / LangChain installed.
* Content types written by this pipeline:
    ``text``          narrative text (TXT, PDF pages, DOCX paragraphs)
    ``table_record``  one generic table row, e.g. a CSV row, with headers
    ``table_summary`` columns + row count + first rows of a table
    ``table_row``     ONLY the legacy STP-schema rows produced by the existing
                      Haryana PDF/XLSX parsers. The structured STP aggregations
                      read this type, so generic tables must never use it.
* Zero extracted records is an error (``UnsupportedDocumentError``), never a
  silent success.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".xlsx", ".csv", ".docx", ".txt"}

# all-MiniLM-L6-v2 truncates around 256 word-pieces (~1000 characters), so
# chunks must stay below that or the tail of the chunk is not embedded at all.
CHUNK_SIZE = 800
CHUNK_OVERLAP = 120

STP_ID_RE = re.compile(r"\b[A-Z]{2,}-STP-\s*\d+\b", re.IGNORECASE)
# Markers that identify rows belonging to the existing STP report schema.
STP_SCHEMA_RE = re.compile(r"\b(?:MLD|STPs?|Compliance|Town)\b", re.IGNORECASE)


class UnsupportedDocumentError(ValueError):
    """The file was readable but produced no indexable content."""


@dataclass
class Chunk:
    page_content: str
    metadata: dict = field(default_factory=dict)


# --------------------------------------------------------------------------
# Text helpers
# --------------------------------------------------------------------------

def read_text_file(path: Path) -> str:
    """Decode a text file without ever raising on odd encodings."""
    data = Path(path).read_bytes()

    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig", errors="replace")
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")

    for encoding in ("utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue

    return data.decode("utf-8", errors="replace")


def _atomize(text: str, size: int, separators=("\n\n", "\n", ". ", " ")) -> list[str]:
    """Break text into pieces no longer than ``size`` at natural boundaries."""
    text = text.strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]

    for index, separator in enumerate(separators):
        parts = text.split(separator)
        if len(parts) == 1:
            continue

        pieces: list[str] = []
        for position, part in enumerate(parts):
            if separator == ". " and position < len(parts) - 1:
                part += "."
            pieces.extend(_atomize(part, size, separators[index + 1:]))
        return pieces

    return [text[i:i + size] for i in range(0, len(text), size)]


def split_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Pack natural-boundary pieces into overlapping chunks of about ``size``."""
    pieces = _atomize(text, size)
    if not pieces:
        return []

    chunks: list[str] = []
    current = ""

    for piece in pieces:
        if current and len(current) + 1 + len(piece) > size:
            chunks.append(current)
            tail = current[-overlap:] if overlap else ""
            # Start the overlap on a word boundary rather than mid-word.
            if tail and " " in tail:
                tail = tail.split(" ", 1)[1]
            current = f"{tail} {piece}".strip() if tail else piece
        else:
            current = f"{current}\n{piece}" if current else piece

    if current:
        chunks.append(current)

    return chunks


# --------------------------------------------------------------------------
# Generic table handling (CSV, DOCX tables, non-STP spreadsheets)
# --------------------------------------------------------------------------

def _clean(value) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def _cell_text(value) -> str:
    """Render a spreadsheet/CSV cell as stable, readable text."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0, 0) else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _clean(value)


def _is_number(text: str) -> bool:
    try:
        float(text.replace(",", ""))
        return True
    except ValueError:
        return False


def _looks_like_header(row: list[str]) -> bool:
    filled = [cell for cell in row if cell]
    return len(filled) >= 2 and not all(_is_number(cell) for cell in filled)


def table_to_records(rows: list[list[str]], base_metadata: dict, label: str) -> list[Chunk]:
    """Turn a table into a summary record plus one labelled record per row."""
    rows = [list(row) for row in rows if any(cell for cell in row)]
    if not rows:
        return []

    if len(rows) > 1 and _looks_like_header(rows[0]):
        headers, body = rows[0], rows[1:]
    else:
        width = max(len(row) for row in rows)
        headers, body = [""] * width, rows

    headers = [header or f"Column {i + 1}" for i, header in enumerate(headers)]

    # Drop columns that are empty in the header and in every body row.
    keep = [
        i for i, header in enumerate(headers)
        if not header.startswith("Column ")
        or any(i < len(row) and row[i] for row in body)
    ]

    def render(row: list[str]) -> str:
        return "\n".join(
            f"{headers[i]}: {row[i]}"
            for i in keep
            if i < len(row) and row[i]
        )

    records: list[Chunk] = []

    summary_lines = [
        f"Table: {label}",
        f"Columns: {', '.join(headers[i] for i in keep)}",
        f"Number of data rows: {len(body)}",
    ]
    for sample_index, row in enumerate(body[:5], start=1):
        summary_lines.append(f"Example row {sample_index}: " + "; ".join(
            f"{headers[i]}={row[i]}" for i in keep if i < len(row) and row[i]
        ))
    records.append(Chunk(
        "\n".join(summary_lines),
        {**base_metadata, "content_type": "table_summary", "row_count": len(body)},
    ))

    for row_number, row in enumerate(body, start=1):
        text = render(row)
        if not text:
            continue
        records.append(Chunk(
            f"Table: {label}\n{text}",
            {**base_metadata, "content_type": "table_record", "row_number": row_number},
        ))

    return records


# --------------------------------------------------------------------------
# Extractors. Each returns un-chunked records.
# --------------------------------------------------------------------------

def extract_txt(path: Path) -> list[Chunk]:
    text = read_text_file(path)
    if not text.strip():
        return []
    return [Chunk(text, {"file_type": "txt", "content_type": "text"})]


def extract_csv(path: Path) -> list[Chunk]:
    text = read_text_file(path)
    if not text.strip():
        return []

    csv.field_size_limit(10_000_000)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    rows = [
        [_cell_text(cell) for cell in row]
        for row in csv.reader(text.splitlines(), dialect)
    ]
    return table_to_records(rows, {"file_type": "csv"}, label=Path(path).name)


def extract_docx(path: Path) -> list[Chunk]:
    from docx import Document as DocxDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = DocxDocument(str(path))
    records: list[Chunk] = []
    paragraphs: list[str] = []
    table_number = 0

    def flush_paragraphs():
        if paragraphs:
            records.append(Chunk(
                "\n".join(paragraphs),
                {"file_type": "docx", "content_type": "text"},
            ))
            paragraphs.clear()

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]

        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if text:
                style = (paragraph.style.name or "") if paragraph.style is not None else ""
                paragraphs.append(f"{text}" if not style.startswith("Heading") else f"\n{text}")

        elif tag == "tbl":
            flush_paragraphs()
            table_number += 1
            table = Table(child, document)
            rows = []
            for row in table.rows:
                cells: list[str] = []
                previous = None
                for cell in row.cells:
                    # python-docx repeats merged cells; keep one copy.
                    if cell._tc is previous:
                        continue
                    previous = cell._tc
                    cells.append(_clean(cell.text))
                rows.append(cells)

            records.extend(table_to_records(
                rows,
                {"file_type": "docx", "table_number": table_number},
                label=f"{Path(path).name} table {table_number}",
            ))

    flush_paragraphs()
    return records


def _looks_like_haryana_stp(path: Path) -> bool:
    """True when the workbook has the STP-report layout the legacy parser needs."""
    from openpyxl import load_workbook

    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(min_row=5, max_row=80, max_col=2, values_only=True):
                heading = row[1] if len(row) > 1 else None
                if heading is not None and STP_ID_RE.search(str(heading)):
                    return True
        return False
    finally:
        workbook.close()


def _generic_xlsx(path: Path) -> list[Chunk]:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    records: list[Chunk] = []

    try:
        for sheet in workbook.worksheets:
            rows = []
            for raw in sheet.iter_rows(values_only=True):
                row = [_cell_text(value) for value in raw]
                while row and not row[-1]:
                    row.pop()
                rows.append(row)
                if len(rows) >= 200_000:
                    break

            # Skip blank / title rows above the real header.
            start = 0
            for index, row in enumerate(rows[:20]):
                following = [r for r in rows[index + 1:index + 4] if any(r)]
                if _looks_like_header(row) and following:
                    start = index
                    break

            records.extend(table_to_records(
                rows[start:],
                {"file_type": "spreadsheet", "sheet": sheet.title, "layout": "generic"},
                label=f"{Path(path).name} / sheet {sheet.title}",
            ))
    finally:
        workbook.close()

    return records


def extract_xlsx(path: Path) -> list[Chunk]:
    """Use the specialised STP parser only when the layout matches."""
    if _looks_like_haryana_stp(path):
        from src.ingestion.load_documents import load_excel_documents

        legacy = load_excel_documents(path)
        if legacy:
            return [
                Chunk(doc.page_content, {**doc.metadata, "layout": "haryana_stp"})
                for doc in legacy
            ]

    return _generic_xlsx(path)


def _ocr_pdf(path: Path) -> list[Chunk]:
    """Best-effort OCR for scanned PDFs; returns [] if OCR is unavailable."""
    try:
        import pdfplumber
        import pytesseract
    except ImportError:
        return []

    records: list[Chunk] = []
    try:
        with pdfplumber.open(path) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                image = page.to_image(resolution=200).original
                text = pytesseract.image_to_string(image)
                if text.strip():
                    records.append(Chunk(text, {
                        "file_type": "pdf",
                        "content_type": "text",
                        "page": page_number,
                        "ocr": True,
                    }))
    except Exception:
        # Missing Tesseract binary, rendering failure, etc.
        return []

    return records


def extract_pdf(path: Path) -> list[Chunk]:
    from src.ingestion.load_documents import load_pdf_documents

    records: list[Chunk] = []
    for doc in load_pdf_documents(path):
        metadata = dict(doc.metadata)

        # Only genuine STP-schema rows keep the legacy "table_row" type.
        # Any other PDF table becomes a generic, searchable table record.
        if metadata.get("content_type") == "table_row" and not STP_SCHEMA_RE.search(doc.page_content):
            metadata["content_type"] = "table_record"

        records.append(Chunk(doc.page_content, metadata))

    return records or _ocr_pdf(path)


def iter_pdf_chunks(path: str | Path):
    """Yield normalized PDF chunks one page at a time for memory-limited hosts."""
    import pdfplumber

    from src.extraction.pdf_table_extractor import (
        extract_page_tables,
    )
    from src.extraction.table_to_text import table_rows_to_text

    path = Path(path).resolve()
    counters: dict[tuple, int] = {}
    found_content = False

    def emit(record: Chunk):
        for part in split_text(record.page_content):
            metadata = {}
            for key, value in record.metadata.items():
                if value is not None:
                    metadata[key] = (
                        value if isinstance(value, (str, int, float, bool)) else str(value)
                    )
            metadata["source"] = str(path)
            metadata["file_name"] = path.name
            metadata.setdefault("content_type", "text")
            group = (
                metadata.get("content_type"), metadata.get("page"),
                metadata.get("sheet"), metadata.get("table_number"),
                metadata.get("row_number"),
            )
            metadata["chunk_index"] = counters.get(group, 0)
            counters[group] = metadata["chunk_index"] + 1
            yield Chunk(part, metadata)

    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            try:
                page_records = []
                tables_on_page = page.find_tables()

                def outside_tables(obj):
                    if obj.get("object_type") != "char":
                        return True
                    for table in tables_on_page:
                        x0, top, x1, bottom = table.bbox
                        if (
                            obj["x0"] >= x0 and obj["x1"] <= x1
                            and obj["top"] >= top and obj["bottom"] <= bottom
                        ):
                            return False
                    return True

                narrative_page = page.filter(outside_tables)
                text = narrative_page.extract_text() or ""
                if text.strip():
                    page_records.append(Chunk(text, {
                        "file_type": "pdf", "content_type": "text", "page": page_number,
                    }))

                for table in extract_page_tables(
                    page, path, page_number, table_objects=tables_on_page
                ):
                    for record in table_rows_to_text(table):
                        metadata = dict(record)
                        if (
                            metadata.get("content_type") == "table_row"
                            and not STP_SCHEMA_RE.search(record["content"])
                        ):
                            metadata["content_type"] = "table_record"
                        page_records.append(Chunk(record["content"], metadata))

                for record in page_records:
                    for chunk in emit(record):
                        found_content = True
                        yield chunk
            finally:
                page.close()

    if not found_content:
        for record in _ocr_pdf(path):
            for chunk in emit(record):
                found_content = True
                yield chunk

    if not found_content:
        raise UnsupportedDocumentError(
            f"No readable content was found in {path.name}. It may be a scanned PDF "
            "without a text layer; OCR (pytesseract + the Tesseract program) is required for scans."
        )


EXTRACTORS = {
    ".txt": extract_txt,
    ".csv": extract_csv,
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
    ".pdf": extract_pdf,
}


# --------------------------------------------------------------------------
# Chunking + metadata normalisation
# --------------------------------------------------------------------------

def chunk_records(records: list[Chunk]) -> list[Chunk]:
    """Split oversized records. Table rows stay whole unless they are huge."""
    chunks: list[Chunk] = []

    for record in records:
        # Legacy STP rows are one-row-per-record by design; never split them.
        if record.metadata.get("content_type") == "table_row":
            chunks.append(record)
            continue

        parts = split_text(record.page_content)
        for part in parts:
            chunks.append(Chunk(part, dict(record.metadata)))

    return chunks


def normalize_metadata(chunks: list[Chunk], path: Path) -> list[Chunk]:
    """Make metadata Chroma-safe and consistent across all file types."""
    absolute = str(Path(path).resolve())
    counters: dict[tuple, int] = {}

    for chunk in chunks:
        metadata = {}
        for key, value in chunk.metadata.items():
            if value is None:
                continue
            metadata[key] = value if isinstance(value, (str, int, float, bool)) else str(value)

        metadata["source"] = absolute
        metadata["file_name"] = Path(path).name
        metadata.setdefault("content_type", "text")

        group = (metadata.get("content_type"), metadata.get("page"), metadata.get("sheet"),
                 metadata.get("table_number"), metadata.get("row_number"))
        metadata["chunk_index"] = counters.get(group, 0)
        counters[group] = metadata["chunk_index"] + 1

        chunk.metadata = metadata

    return chunks


def ingest_file(path: str | Path) -> list[Chunk]:
    """Extract, chunk and normalise one file. Raises if nothing is indexable."""
    path = Path(path)
    extension = path.suffix.lower()

    if not path.is_file():
        raise FileNotFoundError(f"File not found: {path}")
    if extension not in EXTRACTORS:
        raise UnsupportedDocumentError(
            f"Automatic indexing for {extension or 'this file type'} is not supported. "
            "Supported types: PDF, XLSX, CSV, DOCX, TXT."
        )

    chunks = normalize_metadata(chunk_records(EXTRACTORS[extension](path)), path)

    if not chunks:
        hint = (
            " It may be a scanned PDF without a text layer; OCR "
            "(pytesseract + the Tesseract program) is required for scans."
            if extension == ".pdf" else ""
        )
        raise UnsupportedDocumentError(
            f"No readable content was found in {path.name}.{hint}"
        )

    return chunks

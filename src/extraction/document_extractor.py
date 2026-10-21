
from pathlib import Path
import csv

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


def extract_pdf(path: Path) -> list[dict]:
    results = []
    reader = PdfReader(str(path))

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        if text.strip():
            results.append({
                "content": text,
                "source": path.name,
                "file_type": "pdf",
                "page": page_number,
                "content_type": "text",
            })

    return results


def extract_txt(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8", errors="replace")

    return [{
        "content": text,
        "source": path.name,
        "file_type": "txt",
        "page": None,
        "content_type": "text",
    }] if text.strip() else []


def extract_docx(path: Path) -> list[dict]:
    doc = Document(str(path))
    results = []

    for index, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text.strip()

        if text:
            results.append({
                "content": text,
                "source": path.name,
                "file_type": "docx",
                "page": None,
                "paragraph": index,
                "content_type": "text",
            })

    for table_index, table in enumerate(doc.tables, start=1):
        rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in table.rows
        ]
        table_text = "\n".join(rows)

        if table_text.strip():
            results.append({
                "content": table_text,
                "source": path.name,
                "file_type": "docx",
                "page": None,
                "table": table_index,
                "content_type": "table",
            })

    return results


def extract_excel(path: Path) -> list[dict]:
    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )
    results = []

    try:
        for sheet in workbook.worksheets:
            rows = []

            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value).strip()
                    if value is not None else ""
                    for value in row
                ]

                if any(values):
                    rows.append(values)

            if not rows:
                continue

            # Preserve the row and column relationships.
            # Do not treat every row as an independent text chunk.
            table_text = "\n".join(
                " | ".join(row) for row in rows
            )

            results.append({
                "content": table_text,
                "source": path.name,
                "file_type": "excel",
                "sheet": sheet.title,
                "page": None,
                "content_type": "table",
                "row_count": len(rows),
            })
    finally:
        workbook.close()

    return results


def extract_csv(path: Path) -> list[dict]:
    with path.open(
        "r",
        encoding="utf-8-sig",
        errors="replace",
        newline="",
    ) as file:
        rows = list(csv.reader(file))

    if not rows:
        return []

    table_text = "\n".join(
        " | ".join(cell.strip() for cell in row)
        for row in rows
    )

    return [{
        "content": table_text,
        "source": path.name,
        "file_type": "csv",
        "page": None,
        "content_type": "table",
        "row_count": len(rows),
    }]


def extract_document(file_path: str | Path) -> list[dict]:
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"File not found: {path}")

    extension = path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf(path)
    if extension == ".txt":
        return extract_txt(path)
    if extension == ".docx":
        return extract_docx(path)
    if extension in {".xlsx", ".xlsm"}:
        return extract_excel(path)
    if extension == ".csv":
        return extract_csv(path)

    if extension == ".xls":
        raise ValueError(
            "Legacy .xls files need a separate Excel reader."
        )

    raise ValueError(f"Unsupported file type: {extension}")
